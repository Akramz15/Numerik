import io
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_batch_excel(results: list) -> bytes:
    output = io.BytesIO()

    # Create Excel writer using openpyxl engine
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        # 1. Ringkasan Batch Data
        summary_rows = []
        for r in results:
            summary_rows.append(
                {
                    "No": r.get("no"),
                    "Persamaan": r.get("fungsi"),
                    "Batas Bawah (a)": r.get("a"),
                    "Batas Atas (b)": r.get("b"),
                    "Toleransi (Epsilon)": r.get("epsilon"),
                    "Status": "SUKSES" if r.get("status") == "success" else "GAGAL",
                    "Akar Ditemukan": (
                        r.get("root") if r.get("status") == "success" else "-"
                    ),
                    "Jumlah Iterasi": (
                        r.get("iterations_count")
                        if r.get("status") == "success"
                        else "-"
                    ),
                }
            )

        df_summary = pd.DataFrame(summary_rows)
        df_summary.to_excel(writer, sheet_name="Ringkasan Batch", index=False)

        # 2. Detailed Iterations for each successful item
        for r in results:
            if r.get("status") != "success":
                continue

            iter_rows = []
            for it in r.get("iterations", []):
                iter_rows.append(
                    {
                        "Iterasi": it.get("iterasi"),
                        "a": it.get("a"),
                        "b": it.get("b"),
                        "xr (Estimasi Akar)": it.get("xr"),
                        "f(a)": it.get("fa"),
                        "f(b)": it.get("fb"),
                        "f(xr)": it.get("fxr"),
                        "Error": it.get("error"),
                    }
                )

            df_iter = pd.DataFrame(iter_rows)
            sheet_name = f"Soal {r.get('no')}"
            df_iter.to_excel(writer, sheet_name=sheet_name, index=False)

    # Load workbook from memory to apply gorgeous styling
    output.seek(0)
    import openpyxl

    wb = openpyxl.load_workbook(output)

    # Styling definitions
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_regular = Font(name="Calibri", size=11)

    fill_header_summary = PatternFill(
        start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
    )  # Deep navy
    fill_header_detail = PatternFill(
        start_color="0F172A", end_color="0F172A", fill_type="solid"
    )  # Slate-900
    fill_zebra = PatternFill(
        start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
    )  # Alternating rows

    border_thin = Side(style="thin", color="CBD5E1")

    cell_border = Border(
        left=border_thin, right=border_thin, top=border_thin, bottom=border_thin
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    for name in wb.sheetnames:
        ws = wb[name]
        is_summary = name == "Ringkasan Batch"

        # Style Header Row
        ws.row_dimensions[1].height = 28
        header_fill = fill_header_summary if is_summary else fill_header_detail

        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = font_header
            cell.fill = header_fill
            cell.alignment = align_center

        # Style Data Rows
        for r_idx in range(2, ws.max_row + 1):
            ws.row_dimensions[r_idx].height = 20
            use_zebra = r_idx % 2 == 0

            for c_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = font_regular
                cell.border = cell_border

                if use_zebra:
                    cell.fill = fill_zebra

                # Format cell based on columns
                val = cell.value

                # Check alignments and number formats
                if isinstance(val, (int, float)):
                    if not is_summary and c_idx == 1:  # Iteration number
                        cell.alignment = align_center
                    elif is_summary and c_idx in [1, 8]:  # No, Iteration count
                        cell.alignment = align_center
                    else:
                        cell.alignment = align_right
                        # Numbers with float formatting
                        cell.number_format = "0.00000" if not is_summary else "0.0000"
                else:
                    if is_summary and c_idx == 6:  # Status
                        cell.alignment = align_center
                        if val == "SUKSES":
                            cell.font = Font(
                                name="Calibri", size=11, bold=True, color="16A34A"
                            )
                        else:
                            cell.font = Font(
                                name="Calibri", size=11, bold=True, color="DC2626"
                            )
                    else:
                        cell.alignment = align_left

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_idx = col[0].column
            col_letter = get_column_letter(col_idx if col_idx is not None else 1)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Save workbook to final byte array
    final_output = io.BytesIO()
    wb.save(final_output)
    final_output.seek(0)
    return final_output.getvalue()


def generate_single_excel(
    fungsi: str, a: float, b: float, epsilon: float, root: float, iterations: list
) -> bytes:
    output = io.BytesIO()

    # Create Excel writer
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        # Sheet 1: Ringkasan Soal
        summary_rows = [
            {
                "Persamaan Fungsi": fungsi,
                "Batas Bawah (a)": a,
                "Batas Atas (b)": b,
                "Toleransi (Epsilon)": epsilon,
                "Akar Ditemukan": root,
                "Jumlah Iterasi": len(iterations),
            }
        ]
        df_summary = pd.DataFrame(summary_rows)
        df_summary.to_excel(writer, sheet_name="Ringkasan", index=False)

        # Sheet 2: Tabel Iterasi
        iter_rows = []
        for it in iterations:
            iter_rows.append(
                {
                    "Iterasi": it.get("iterasi"),
                    "a": it.get("a"),
                    "b": it.get("b"),
                    "xr (Estimasi Akar)": it.get("xr"),
                    "f(a)": it.get("fa"),
                    "f(b)": it.get("fb"),
                    "f(xr)": it.get("fxr"),
                    "Error": it.get("error"),
                }
            )
        df_iter = pd.DataFrame(iter_rows)
        df_iter.to_excel(writer, sheet_name="Tabel Iterasi", index=False)

    output.seek(0)
    import openpyxl

    wb = openpyxl.load_workbook(output)

    # Styling definitions
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_regular = Font(name="Calibri", size=11)

    fill_header_summary = PatternFill(
        start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
    )  # Deep navy
    fill_header_detail = PatternFill(
        start_color="0F172A", end_color="0F172A", fill_type="solid"
    )  # Slate-900
    fill_zebra = PatternFill(
        start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
    )  # Alternating rows

    border_thin = Side(style="thin", color="CBD5E1")
    cell_border = Border(
        left=border_thin, right=border_thin, top=border_thin, bottom=border_thin
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    for name in wb.sheetnames:
        ws = wb[name]
        is_summary = name == "Ringkasan"

        # Style Header
        ws.row_dimensions[1].height = 28
        header_fill = fill_header_summary if is_summary else fill_header_detail

        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = font_header
            cell.fill = header_fill
            cell.alignment = align_center

        # Style Data Rows
        for r_idx in range(2, ws.max_row + 1):
            ws.row_dimensions[r_idx].height = 20
            use_zebra = r_idx % 2 == 0

            for c_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.font = font_regular
                cell.border = cell_border

                if use_zebra:
                    cell.fill = fill_zebra

                val = cell.value

                if isinstance(val, (int, float)):
                    if not is_summary and c_idx == 1:  # Iteration number
                        cell.alignment = align_center
                    else:
                        cell.alignment = align_right
                        cell.number_format = "0.00000"
                else:
                    cell.alignment = align_left

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_id = col[0].column
            col_letter = get_column_letter(col_id if col_id is not None else 1)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    final_output = io.BytesIO()
    wb.save(final_output)
    final_output.seek(0)
    return final_output.getvalue()
